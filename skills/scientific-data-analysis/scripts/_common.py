from __future__ import annotations

import hashlib
import json
import math
from datetime import date, datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd

SKILL_VERSION = "1.6.1"
ANALYSIS_PLAN_SCHEMA_VERSION = "1.4"
POWER_PLAN_SCHEMA_VERSION = "1.1"




def categorical_level_key(value: Any):
    """Stable categorical identity preserving strings versus numeric scalars.

    Numeric 1 and 1.0 are treated as the same data level, while string "1"
    remains distinct. Numpy scalar wrappers are normalized to Python scalars.
    """
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, bool):
        return ("bool", bool(value))
    if isinstance(value, str):
        return ("str", value)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        v = float(value)
        if not math.isfinite(v):
            return ("nonfinite", repr(value))
        return ("num", v)
    if value is None:
        return ("none", None)
    return (type(value).__name__, repr(value))




def typed_scalar_key(value: Any):
    """Identity key for tabular scalars, treating all missing values alike.

    This extends categorical_level_key to row/identifier auditing. Numeric 1
    and 1.0 intentionally share an identity, while boolean True and numeric 1
    remain distinct. Missing values map to one explicit missing identity.
    """
    try:
        if pd.isna(value):
            return ("missing", None)
    except Exception:
        pass
    return categorical_level_key(value)


def typed_duplicate_mask(series: pd.Series, keep=False) -> pd.Series:
    """Duplicate mask using typed scalar identity instead of pandas equality."""
    keys = pd.Series([typed_scalar_key(v) for v in series.tolist()], index=series.index, dtype=object)
    return keys.duplicated(keep=keep)


def typed_nunique(series: pd.Series, dropna: bool = True) -> int:
    """Count unique scalar identities without bool/numeric equality collisions."""
    keys = set()
    for v in series.tolist():
        key = typed_scalar_key(v)
        if dropna and key[0] == "missing":
            continue
        keys.add(key)
    return len(keys)


def typed_row_duplicate_mask(df: pd.DataFrame, keep="first") -> pd.Series:
    """Exact-row duplicate mask under the skill's typed scalar identity."""
    row_keys = [tuple(typed_scalar_key(v) for v in row) for row in df.itertuples(index=False, name=None)]
    return pd.Series(row_keys, index=df.index, dtype=object).duplicated(keep=keep)

def categorical_level_mask(series: pd.Series, level: Any) -> pd.Series:
    """Boolean mask using typed categorical identity rather than Python/pandas equality.

    This prevents equality-colliding encodings such as ``True`` and ``1`` from
    being silently merged, while still treating numeric ``1`` and ``1.0`` as
    the same scalar level. Missing values never match a declared level.
    """
    target = categorical_level_key(level)
    def _matches(v):
        try:
            if pd.isna(v):
                return False
        except Exception:
            return False
        return categorical_level_key(v) == target
    return series.map(_matches).astype(bool)


def categorical_levels_mask(series: pd.Series, levels) -> pd.Series:
    """Union mask for multiple declared levels using typed categorical identity."""
    levels = list(levels)
    if not levels:
        return pd.Series(False, index=series.index, dtype=bool)
    out = pd.Series(False, index=series.index, dtype=bool)
    for level in levels:
        out |= categorical_level_mask(series, level)
    return out


def assert_no_equality_colliding_levels(series: pd.Series, *, label: str = "categorical variable") -> None:
    """Reject distinct typed levels that Python/pandas equality would merge.

    Numeric 1 and 1.0 intentionally share one typed identity. Distinct typed
    identities such as True vs 1 are unsafe for pandas grouping/crosstab
    because they compare equal and can be silently collapsed before later
    typed-identity checks run.
    """
    levels = []
    seen = set()
    for value in series.tolist():
        try:
            if pd.isna(value):
                continue
        except Exception:
            pass
        key = categorical_level_key(value)
        if key in seen:
            continue
        seen.add(key)
        levels.append(value)
    for i in range(len(levels)):
        for j in range(i + 1, len(levels)):
            a, b = levels[i], levels[j]
            if categorical_level_key(a) == categorical_level_key(b):
                continue
            try:
                equal = bool(a == b)
            except Exception:
                equal = False
            if equal:
                raise ValueError(
                    f"{label} contains distinct typed categorical levels that Python/pandas equality would merge: "
                    f"{jsonable_level(a)!r} ({categorical_level_key(a)[0]}) vs "
                    f"{jsonable_level(b)!r} ({categorical_level_key(b)[0]}); normalize the categorical encoding before analysis"
                )


def jsonable_level(value: Any):
    """Return a JSON-safe scalar while preserving categorical level type."""
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, bool):
        return bool(value)
    if isinstance(value, str):
        return value
    if isinstance(value, int) and not isinstance(value, bool):
        return int(value)
    if isinstance(value, float):
        return float(value) if math.isfinite(float(value)) else None
    return str(value)



def level_display_labels(values):
    """Human-readable categorical labels that disambiguate type collisions."""
    vals = [jsonable_level(v) for v in values]
    base = [str(v) for v in vals]
    counts = {b: base.count(b) for b in set(base)}
    out = []
    for v, b in zip(vals, base):
        if counts[b] <= 1:
            out.append(b)
            continue
        if isinstance(v, str):
            kind = "string"
        elif isinstance(v, bool):
            kind = "boolean"
        elif isinstance(v, (int, float)) and not isinstance(v, bool):
            kind = "number"
        else:
            kind = type(v).__name__
        out.append(f"{b} [{kind}]")
    return out

def _read_xlsx_preserve_scalars(path: str | Path) -> pd.DataFrame:
    """Read the first XLSX/XLSM sheet without pandas scalar-type coercion.

    pandas.read_excel may coerce a column containing Excel boolean TRUE and
    numeric 1 into one homogeneous dtype before the skill can apply its typed
    identity rules. Reading native cell values through openpyxl preserves the
    distinction long enough for cleaning/validation to make an explicit choice.
    """
    from openpyxl import load_workbook

    wb = load_workbook(Path(path), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    if not rows:
        return pd.DataFrame()
    width = max((i + 1 for row in rows for i, v in enumerate(row) if v is not None), default=0)
    if width == 0:
        return pd.DataFrame()
    rows = [(row + [None] * width)[:width] for row in rows]
    header = rows[0]
    if any(h is None or str(h).strip() == "" for h in header):
        raise ValueError("Excel input contains blank column headers; rename columns before analysis")
    header = [str(h) for h in header]
    if len(header) != len(set(header)):
        raise ValueError("Excel input contains duplicate column headers; rename columns before analysis")
    return pd.DataFrame(rows[1:], columns=header, dtype=object)


def read_table(path: str | Path) -> pd.DataFrame:
    p = Path(path)
    suf = p.suffix.lower()
    if suf == ".csv":
        return pd.read_csv(p)
    if suf in {".tsv", ".txt"}:
        return pd.read_csv(p, sep="\t")
    if suf in {".xlsx", ".xlsm"}:
        return _read_xlsx_preserve_scalars(p)
    if suf == ".xls":
        return pd.read_excel(p)
    if suf in {".parquet", ".pq"}:
        return pd.read_parquet(p)
    if suf in {".jsonl", ".ndjson"}:
        # dtype=False is required to keep JSON booleans, numbers and strings
        # distinct in mixed categorical/identifier columns.
        return pd.read_json(p, lines=True, dtype=False, precise_float=True)
    raise ValueError(f"Unsupported table format: {p.suffix}")


def read_artifact_csv(path: str | Path) -> pd.DataFrame:
    """Read generated result/sensitivity CSV without re-inferring JSON text.

    pandas otherwise parses a cell containing the JSON scalar text ``true`` as
    a boolean, which breaks deterministic reconciliation against freshly
    generated JSON metadata. Columns ending in ``_json`` are artifact text and
    must remain text.
    """
    p = Path(path)
    header = pd.read_csv(p, nrows=0)
    dtype = {c: "string" for c in header.columns if str(c).endswith("_json")}
    return pd.read_csv(p, dtype=dtype or None)


def write_table(df: pd.DataFrame, path: str | Path) -> None:
    """Write an analysis-ready table using an extension-defined format.

    JSONL is the type-preserving interchange format for mixed scalar identity
    (e.g. boolean true versus numeric 1 versus string "1"). CSV/TSV remain
    convenient defaults when a round-trip fidelity check confirms that no
    plan-relevant scalar identity changes.
    """
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    suf = p.suffix.lower()
    if suf == ".csv":
        df.to_csv(p, index=False)
        return
    if suf in {".tsv", ".txt"}:
        df.to_csv(p, sep="\t", index=False)
        return
    if suf in {".jsonl", ".ndjson"}:
        with p.open("w", encoding="utf-8") as f:
            for record in df.to_dict(orient="records"):
                f.write(json.dumps(_json_sanitize(record), ensure_ascii=False, allow_nan=False) + "\n")
        return
    raise ValueError(
        f"Unsupported analysis-ready output format: {p.suffix}. "
        "Use .csv/.tsv for ordinary data or .jsonl when scalar type identity must be preserved."
    )


def table_roundtrip_identity_mismatches(
    original: pd.DataFrame, reloaded: pd.DataFrame, columns, *, max_examples: int = 10
):
    """Return examples where a table serialization changed typed scalar identity."""
    cols = [c for c in columns if c in original.columns and c in reloaded.columns]
    mismatches = []
    if len(original) != len(reloaded):
        return [{"reason": "row_count_changed", "before": int(len(original)), "after": int(len(reloaded))}]
    for c in cols:
        before = original[c].tolist()
        after = reloaded[c].tolist()
        for pos, (a, b) in enumerate(zip(before, after)):
            if typed_scalar_key(a) != typed_scalar_key(b):
                mismatches.append({
                    "column": c, "row_position": int(pos + 1),
                    "before": jsonable_level(a) if typed_scalar_key(a)[0] != "missing" else None,
                    "before_type": typed_scalar_key(a)[0],
                    "after": jsonable_level(b) if typed_scalar_key(b)[0] != "missing" else None,
                    "after_type": typed_scalar_key(b)[0],
                })
                if len(mismatches) >= max_examples:
                    return mismatches
    return mismatches


def _json_sanitize(x: Any):
    """Convert numpy/pandas values and non-finite numbers to portable JSON values."""
    if isinstance(x, dict):
        return {str(k): _json_sanitize(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [_json_sanitize(v) for v in x]
    if isinstance(x, np.ndarray):
        return [_json_sanitize(v) for v in x.tolist()]
    if isinstance(x, (np.integer,)):
        return int(x)
    if isinstance(x, (np.floating, float)):
        v = float(x)
        return v if math.isfinite(v) else None
    if isinstance(x, Path):
        return str(x)
    if isinstance(x, (pd.Timestamp, datetime, date)):
        return x.isoformat()
    try:
        if pd.isna(x):
            return None
    except Exception:
        pass
    return x


def write_json(obj: Any, path: str | Path) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    with p.open("w", encoding="utf-8") as f:
        json.dump(_json_sanitize(obj), f, ensure_ascii=False, indent=2, allow_nan=False)


def read_json(path: str | Path) -> dict[str, Any]:
    with Path(path).open("r", encoding="utf-8") as f:
        return json.load(f)


def sha256_file(path: str | Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def stable_seed(base_seed: int, label: str) -> int:
    """Derive an order-independent deterministic uint32 seed from a plan seed + label."""
    payload = f"{int(base_seed)}::{label}".encode("utf-8")
    return int.from_bytes(hashlib.sha256(payload).digest()[:4], "big", signed=False)


def _json_default(x: Any):
    if isinstance(x, (np.integer,)):
        return int(x)
    if isinstance(x, (np.floating,)):
        return float(x)
    if isinstance(x, np.ndarray):
        return x.tolist()
    if isinstance(x, Path):
        return str(x)
    if isinstance(x, (pd.Timestamp, datetime, date)):
        return x.isoformat()
    if pd.isna(x):
        return None
    raise TypeError(type(x).__name__)


def safe_float(x: Any) -> float | None:
    try:
        v = float(x)
        return v if math.isfinite(v) else None
    except Exception:
        return None


def ci_level_to_alpha(level: float) -> float:
    if not (0 < level < 1):
        raise ValueError("confidence_level must be between 0 and 1")
    return 1.0 - level


def bootstrap_ci(stat_fn, arrays, *, confidence_level=0.95, seed=0, n_resamples=2000, paired=False):
    """Deterministic percentile bootstrap CI.

    Invalid resamples (for example, constant-input correlation resamples) are
    skipped rather than aborting the whole analysis. If too few valid
    resamples remain, the interval is returned as not estimable.
    """
    rng = np.random.default_rng(seed)
    arrays = [np.asarray(a) for a in arrays]
    if any(len(a) == 0 for a in arrays):
        return (None, None)
    stats_out = []
    if paired:
        n = len(arrays[0])
        if any(len(a) != n for a in arrays):
            raise ValueError("Paired bootstrap arrays must have equal length")
        for _ in range(n_resamples):
            idx = rng.integers(0, n, size=n)
            try:
                value = stat_fn(*[a[idx] for a in arrays])
            except Exception:
                continue
            if np.isfinite(value):
                stats_out.append(value)
    else:
        for _ in range(n_resamples):
            resampled = []
            for a in arrays:
                idx = rng.integers(0, len(a), size=len(a))
                resampled.append(a[idx])
            try:
                value = stat_fn(*resampled)
            except Exception:
                continue
            if np.isfinite(value):
                stats_out.append(value)
    stats_out = np.asarray(stats_out, dtype=float)
    if stats_out.size < max(100, n_resamples // 20):
        return (None, None)
    alpha = 1 - confidence_level
    return (
        float(np.quantile(stats_out, alpha / 2)),
        float(np.quantile(stats_out, 1 - alpha / 2)),
    )
